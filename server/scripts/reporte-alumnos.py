#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Reporte de Alumnos (Corporativo)
--------------------------------
- ESTUDIANTES con perfilCompletado=1
- Último intento por quiz (BAI, BDI2...) por alumno
- Grupo en MAYÚSCULAS
- Hojas: README, Resumen, {INI}_{GRUPO}
- Estilo corporativo (Montserrat, colores, zebra, tablas)

Uso:
  python reporte-alumnos.py \
    --host 127.0.0.1 --port 3306 \
    --user root --password "mon123" \
    --database sistema_educativo --outfile reporte_alumnos_corporativo.xlsx
"""
import argparse
import sys
import re, uuid
from datetime import datetime

import pandas as pd
import mysql.connector as mc
from slugify import slugify

# ===================== PALETA CORPORATIVA =====================
COLORS = {
    "cerulean": "019fd2",       # headers principales
    "picton": "48b0f7",         # headers secundarios / títulos
    "outer_space": "2b333c",    # texto oscuro
    "wild_sand": "f7f7f7",      # fondo claro
    "empress": "7c777a",        # texto secundario
    "white": "ffffff",
    "zebra": "f0f7fb",          # filas alternas
}

DEFAULT_OUTFILE = "reporte_alumnos_corporativo.xlsx"

def _sanitize_table_name(name: str) -> str:
    """
    Excel table displayName:
      - sin espacios ni caracteres raros
      - debe iniciar con letra o guion bajo
      - debe ser único en el libro
    """
    base = re.sub(r'[^A-Za-z0-9_]', '_', str(name))
    if not re.match(r'^[A-Za-z_]', base):
        base = f"T_{base}"
    unique = uuid.uuid4().hex[:6]
    return (f"{base}_{unique}")[:31]

# ===================== CONEXIÓN =====================
def connect_mysql(host, port, user, password, database):
    try:
        conn = mc.connect(
            host=host, port=port, user=user, password=password,
            database=database, autocommit=False
        )
        return conn
    except mc.Error as e:
        print(f"[ERROR] No se pudo conectar a MySQL: {e}")
        sys.exit(1)

# ===================== QUERIES =====================
def fetch_data(conn):
    users_sql = """
        SELECT
            u.id                         AS usuarioId,
            u.carrera,
            u.email,
            u.nombreCompleto,
            u.matricula,
            u.genero,
            u.telefono,
            u.semestre,
            UPPER(COALESCE(u.grupo, '')) AS grupo,
            u.turno
        FROM usuarios u
        WHERE u.rol = 'ESTUDIANTE'
          AND u.perfilCompletado = 1
    """
    quizzes_sql = """
        SELECT q.id AS quizId, q.codigo, q.titulo
        FROM quizzes q
        WHERE q.activo = 1
    """
    latest_attempts_sql = """
        WITH ultimos AS (
            SELECT
                rq.usuarioId,
                rq.quizId,
                rq.puntajeTotal,
                rq.severidad,
                rq.completado,
                rq.fechaEnvio,
                ROW_NUMBER() OVER (PARTITION BY rq.usuarioId, rq.quizId ORDER BY rq.fechaEnvio DESC) AS rn
            FROM respuestas_quiz rq
        )
        SELECT
            usuarioId,
            quizId,
            puntajeTotal,
            severidad,
            completado,
            fechaEnvio
        FROM ultimos
        WHERE rn = 1
    """
    users_df   = pd.read_sql(users_sql, conn)
    quizzes_df = pd.read_sql(quizzes_sql, conn)
    latest_df  = pd.read_sql(latest_attempts_sql, conn)
    return users_df, quizzes_df, latest_df

def pivot_latest(users_df, quizzes_df, latest_df):
    if users_df.empty:
        return users_df.copy()

    code_map = quizzes_df.set_index('quizId')['codigo'].to_dict()
    latest_df = latest_df.copy()
    latest_df['codigo'] = latest_df['quizId'].map(code_map)
    latest_df = latest_df[latest_df['codigo'].notna()].copy()

    # Larga → pivot por métrica
    long_rows = []
    for col, suffix in [('puntajeTotal','puntaje'), ('severidad','severidad'),
                        ('completado','completado'), ('fechaEnvio','fecha')]:
        tmp = latest_df[['usuarioId','codigo',col]].copy()
        tmp['metric'] = suffix
        tmp.rename(columns={col: 'valor'}, inplace=True)
        long_rows.append(tmp)

    long_df = pd.concat(long_rows, ignore_index=True) if long_rows else pd.DataFrame(
        columns=['usuarioId','codigo','metric','valor']
    )
    long_df['colname'] = long_df['codigo'].astype(str) + "_" + long_df['metric'].astype(str)
    pivot_df = long_df.pivot_table(
        index='usuarioId', columns='colname', values='valor', aggfunc='first'
    ).reset_index()

    full_df = users_df.merge(pivot_df, how='left', on='usuarioId')

    base_cols = [
        'usuarioId','matricula','nombreCompleto','email','telefono',
        'genero','semestre','turno','carrera','grupo'
    ]
    other_cols = [c for c in full_df.columns if c not in base_cols]
    ordered = base_cols + sorted(other_cols)
    full_df = full_df.reindex(columns=ordered)

    full_df['grupo'] = full_df['grupo'].fillna('').astype(str).str.upper()
    return full_df

# ===================== ESTILO EXCEL =====================
def auto_fit_columns(ws, df, start_col=1, max_width=48):
    from openpyxl.utils import get_column_letter
    if df is None or df.empty:
        return
    for i, col in enumerate(df.columns, start_col):
        series = df[col].astype(str)
        max_len = max([len(str(col))] + [len(s) for s in series.tolist()]) if not series.empty else len(str(col))
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, max_len + 2), max_width)

# Dejamos esta función inofensiva para no crear tablas estructuradas (evitar corrupción del archivo)
def apply_table_style(ws, df, table_name="Tabla1", start_row=3, start_col=1):
    return

def paint_header(ws, df, start_row=3, start_col=1):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    th_font = Font(name="Montserrat", bold=True, color=COLORS["white"])
    th_fill = PatternFill("solid", fgColor=COLORS["cerulean"])
    align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style="thin", color=COLORS["cerulean"]),
        right=Side(style="thin", color=COLORS["cerulean"]),
        top=Side(style="thin", color=COLORS["cerulean"]),
        bottom=Side(style="thin", color=COLORS["cerulean"])
    )

    for j in range(df.shape[1]):
        cell = ws.cell(row=start_row, column=start_col + j)
        cell.font = th_font
        cell.fill = th_fill
        cell.alignment = align
        cell.border = border
        ws.row_dimensions[start_row].height = 22

def zebra_rows(ws, df, start_row=4, start_col=1):
    from openpyxl.styles import PatternFill
    fill = PatternFill("solid", fgColor=COLORS["zebra"])
    for r in range(df.shape[0]):
        if r % 2 == 0:
            for c in range(df.shape[1]):
                ws.cell(row=start_row + r, column=start_col + c).fill = fill

def write_title(ws, title, sub=None):
    from openpyxl.styles import Font, Alignment
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ws.max_column, 1))
    ws["A1"] = title
    ws["A1"].font = Font(name="Montserrat", bold=True, size=16, color=COLORS["outer_space"])
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ws.max_column, 1))
        ws["A2"] = sub
        ws["A2"].font = Font(name="Montserrat", size=11, color=COLORS["empress"])

def write_df(ws, df, title, subtitle=None, table_base="Tabla", freeze_header=True):
    from openpyxl.utils.dataframe import dataframe_to_rows
    from openpyxl.utils import get_column_letter

    # Escribe título
    write_title(ws, title, subtitle)

    # Encabezados y datos a partir de fila 3
    start_row = 3
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start_row):
        for c_idx, value in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=value)

    # Estilo encabezado + zebra
    paint_header(ws, df, start_row=start_row, start_col=1)
    zebra_rows(ws, df, start_row=start_row + 1, start_col=1)

    # Definir rango real de datos para autofiltro (solo encabezado + filas)
    if df.shape[0] > 0 and df.shape[1] > 0:
        last_row = start_row + df.shape[0]
        last_col_letter = get_column_letter(df.shape[1])
        ws.auto_filter.ref = f"A{start_row}:{last_col_letter}{last_row}"

    # Freeze panes justo debajo del header
    if freeze_header:
        ws.freeze_panes = ws["A4"]

    # Auto-ajuste de columnas
    auto_fit_columns(ws, df, start_col=1)

    # No agregamos Table() para evitar que Excel marque el archivo como dañado
    # apply_table_style(ws, df, table_name=f"{table_base}", start_row=start_row, start_col=1)

# ===================== ESCRITURA EXCEL =====================
def safe_sheet_name(name: str) -> str:
    name = re.sub(r'[:\\/?*\[\]]', '-', name)
    return name[:31] if len(name) > 31 else (name or "Hoja")

def initials_from_carrera(carrera: str) -> str:
    if not carrera or str(carrera).strip().upper() in ('', 'NULL'):
        return "SIN_CARRERA"
    words = re.split(r'\s+', str(carrera).strip())
    ini = ''.join([w[0] for w in words if w and w[0].isalpha()]).upper()
    return ini or "CARR"

def write_excel(full_df: pd.DataFrame, outfile: str):
    import openpyxl

    with pd.ExcelWriter(outfile, engine='openpyxl', datetime_format="yyyy-mm-dd hh:mm:ss") as writer:
        # README
        readme = pd.DataFrame({
            "Información": [
                "REPORTE DE ALUMNOS — CORPORATIVO",
                "",
                "Hojas:",
                " - README: esta información.",
                " - Resumen: conteo de alumnos por Carrera y Grupo.",
                " - {INI}_{GRUPO}: detalle con últimas métricas de cada quiz por alumno.",
                "",
                "Notas:",
                " - 'grupo' normalizado a MAYÚSCULAS; si falta, aparece SIN GRUPO.",
                " - Por cada quiz aparecen columnas: _puntaje, _severidad, _completado, _fecha.",
                " - Si alguien no ha intentado un quiz, esas columnas se verán vacías."
            ]
        })
        readme.to_excel(writer, index=False, sheet_name="README")
        ws = writer.sheets["README"]
        write_df(ws, readme, title="README", subtitle="Guía rápida del reporte", table_base="ReadmeTable")

        if full_df.empty:
            return

        # Resumen
        resumen = (full_df.assign(grupo=full_df['grupo'].fillna('').str.upper())
                            .groupby(['carrera','grupo'], dropna=False)
                            .size()
                            .reset_index(name='alumnos'))
        resumen = resumen.sort_values(['carrera','grupo'], kind='stable')
        resumen.to_excel(writer, index=False, sheet_name="Resumen")
        ws_res = writer.sheets["Resumen"]
        write_df(
            ws_res, resumen,
            title="Resumen por Carrera y Grupo",
            subtitle=f"Generado: {datetime.now():%Y-%m-%d %H:%M}",
            table_base="ResumenTable"
        )

        # Detalle por grupo
        full_df['carrera'] = full_df['carrera'].fillna('SIN CARRERA')
        full_df['grupo'] = full_df['grupo'].replace({'': 'SIN GRUPO'})

        for (carrera, grupo), df_grp in full_df.groupby(['carrera','grupo'], dropna=False):
            ini = initials_from_carrera(carrera)
            sheet_name = safe_sheet_name(f"{ini}_{grupo}")

            base_cols = [
                'usuarioId','matricula','nombreCompleto','email','telefono',
                'genero','semestre','turno','carrera','grupo'
            ]
            other_cols = [c for c in df_grp.columns if c not in base_cols]
            ordered = base_cols + sorted(other_cols)
            df_out = df_grp.reindex(columns=ordered).sort_values(
                ['grupo','semestre','nombreCompleto'], kind='stable'
            )

            df_out.to_excel(writer, index=False, sheet_name=sheet_name)
            ws_det = writer.sheets[sheet_name]
            write_df(
                ws_det, df_out,
                title=f"{carrera} — Grupo {grupo}",
                subtitle="Detalle de alumnos y última actividad por quiz",
                table_base=f"Det_{ini}_{grupo}"
            )

# ===================== MAIN =====================
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--host', default='127.0.0.1')
    ap.add_argument('--port', type=int, default=3306)
    ap.add_argument('--user', default='root')
    ap.add_argument('--password', default='mon123')
    ap.add_argument('--database', default='sistema_educativo')
    ap.add_argument('--outfile', default=DEFAULT_OUTFILE)
    args = ap.parse_args()

    conn = connect_mysql(args.host, args.port, args.user, args.password, args.database)
    try:
        users_df, quizzes_df, latest_df = fetch_data(conn)
    finally:
        conn.close()

    full_df = pivot_latest(users_df, quizzes_df, latest_df)
    write_excel(full_df, args.outfile)
    print(f"[OK] Reporte generado: {args.outfile}")

if __name__ == "__main__":
    main()
